import openpyxl
from openpyxl import load_workbook
import json

# Load the Excel file
excel_path = r'C:\Users\kurok\Downloads\G_ADC\ADC_08\テスト＿卒論.xlsx'

wb = load_workbook(excel_path)

output = []
output.append(f"Available sheets: {wb.sheetnames}\n")

# Examine each sheet
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    output.append(f"\n{'='*80}")
    output.append(f"Sheet: {sheet_name}")
    output.append(f"{'='*80}")
    output.append(f"Dimensions: {ws.max_row} rows x {ws.max_column} columns\n")
    
    # Display the content with row numbers
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(50, ws.max_row), values_only=True), start=1):
        cells = []
        for cell_idx, cell in enumerate(row, start=1):
            if cell is None or str(cell).strip() == "":
                cells.append(f"[空{cell_idx}]")
            else:
                cells.append(f"{cell}")
        output.append(f"R{row_idx:02d}: " + " | ".join(cells))

# Write to file
with open('excel_structure.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(output))

print("Analysis saved to excel_structure.txt")
