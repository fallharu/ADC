import openpyxl
from openpyxl import load_workbook

# Load the Excel file
excel_path = r'C:\Users\kurok\Downloads\G_ADC\ADC_08\テスト＿卒論.xlsx'

wb = load_workbook(excel_path)
print(f"Available sheets: {wb.sheetnames}\n")

# Examine each sheet
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"{'='*80}")
    print(f"Sheet: {sheet_name}")
    print(f"{'='*80}")
    print(f"Dimensions: {ws.max_row} rows x {ws.max_column} columns\n")
    
    # Display the content with row numbers
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(30, ws.max_row), values_only=True), start=1):
        cells = []
        for cell_idx, cell in enumerate(row, start=1):
            if cell is None:
                cells.append(f"Col{cell_idx}:[空]")
            else:
                cells.append(f"Col{cell_idx}:{str(cell)[:30]}")
        print(f"Row{row_idx}: " + " | ".join(cells))
    print("\n")
