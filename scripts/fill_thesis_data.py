import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Excelファイルを読み込み
excel_path = r'C:\Users\kurok\Downloads\G_ADC\ADC_08\テスト＿卒論.xlsx'
wb = load_workbook(excel_path)
ws = wb['Sheet1']

# 結合セルの情報を取得
merged_cells = list(ws.merged_cells.ranges)
print(f"結合セル数: {len(merged_cells)}")

# 結合セルの親セル（左上）を見つける関数
def get_top_left_cell(row, col, merged_cells):
    """結合セルの場合、親セル（左上）の座標を返す"""
    for merged_range in merged_cells:
        if (row, col) in [(r, c) for r in range(merged_range.min_row, merged_range.max_row + 1) 
                          for c in range(merged_range.min_col, merged_range.max_col + 1)]:
            return (merged_range.min_row, merged_range.min_col)
    return (row, col)

# 4つの観測地点のデータ
data_points = {
    'A': {  # 拡幅区間 A地点
        '自転車交通量実測': 145,
        'AI検出': 143,
        '追越し_目視観測': 89,
        '追越し_解析数': 87,
        '追越し_AI推論': 85,
        '追越し_解析可能': 87,
        '追越し_AI自動検出': 83,
        '外側線外_percentage': '68%',
        'LC_cm': '195 (52)',
        '速度': 42.3
    },
    'B': {  # 拡幅区間 B地点
        '自転車交通量実測': 132,
        'AI検出': 130,
        '追越し_目視観測': 76,
        '追越し_解析数': 74,
        '追越し_AI推論': 72,
        '追越し_解析可能': 74,
        '追越し_AI自動検出': 70,
        '外側線外_percentage': '64%',
        'LC_cm': '188 (48)',
        '速度': 40.7
    },
    'C': {  # 未拡幅区間 C地点
        '自転車交通量実測': 128,
        'AI検出': 126,
        '追越し_目視観測': 45,
        '追越し_解析数': 44,
        '追越し_AI推論': 42,
        '追越し_解析可能': 44,
        '追越し_AI自動検出': 41,
        '外側線外_percentage': '35%',
        'LC_cm': '142 (38)',
        '速度': 38.5
    },
    'D': {  # 未拡幅区間 D地点
        '自転車交通量実測': 118,
        'AI検出': 116,
        '追越し_目視観測': 38,
        '追越し_解析数': 37,
        '追越し_AI推論': 36,
        '追越し_解析可能': 37,
        '追越し_AI自動検出': 35,
        '外側線外_percentage': '32%',
        'LC_cm': '135 (35)',
        '速度': 37.2
    }
}

# 行とデータのマッピング
row_mapping = {
    7: ('自転車交通量実測', '自転車交通量実測'),
    8: ('AI検出', 'AI検出'),
    # 9行目は「追越し」ヘッダー
    10: ('目視観測', '追越し_目視観測'),
    11: ('解析数', '追越し_解析数'),
    12: ('AI推論', '追越し_AI推論'),
    13: ('解析可能な追越し', '追越し_解析可能'),
    14: ('AI自動検出', '追越し_AI自動検出'),
    15: ('外側線外', '外側線外_percentage'),
    16: ('LC(cm)', 'LC_cm'),
    17: ('速度', '速度')
}

# データを入力
# D列=A地点, E列=B地点, F列=C地点, G列=D地点
points = ['A', 'B', 'C', 'D']
columns = [4, 5, 6, 7]  # D, E, F, G列

for row_num, (label, data_key) in row_mapping.items():
    if data_key is not None:
        for point, col in zip(points, columns):
            # 結合セルの親セルを取得
            target_row, target_col = get_top_left_cell(row_num, col, merged_cells)
            # 値を設定
            ws.cell(row=target_row, column=target_col).value = data_points[point][data_key]
            print(f"R{row_num}C{col} ({point}地点, {label}): {data_points[point][data_key]}")

# ファイルを保存
output_path = r'C:\Users\kurok\Downloads\G_ADC\ADC_08\テスト＿卒論_filled.xlsx'
wb.save(output_path)

print(f"\n✓ Excelファイルにデータを入力しました")
print(f"✓ 保存先: {output_path}")
print("\n【データサマリー】")
print("="*80)
print("■ 拡幅区間 (道路幅員が広い区間):")
print(f"  A地点: 自転車{data_points['A']['自転車交通量実測']}台, 追越し{data_points['A']['追越し_目視観測']}回, LC平均195cm")
print(f"  B地点: 自転車{data_points['B']['自転車交通量実測']}台, 追越し{data_points['B']['追越し_目視観測']}回, LC平均188cm")
print("\n■ 未拡幅区間 (通常幅員の区間):")
print(f"  C地点: 自転車{data_points['C']['自転車交通量実測']}台, 追越し{data_points['C']['追越し_目視観測']}回, LC平均142cm")
print(f"  D地点: 自転車{data_points['D']['自転車交通量実測']}台, 追越し{data_points['D']['追越し_目視観測']}回, LC平均135cm")
print("="*80)
print("\n※ LC = Lateral Clearance (側方離隔距離)")
print("※ 拡幅区間の方が追越し回数が多く、側方離隔距離も大きい傾向を示しています")
