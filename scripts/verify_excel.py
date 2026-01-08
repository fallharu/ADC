import openpyxl
from openpyxl import load_workbook

# 入力後のファイルを読み込んで確認
excel_path = r'C:\Users\kurok\Downloads\G_ADC\ADC_08\テスト＿卒論_filled.xlsx'
wb = load_workbook(excel_path)
ws = wb['Sheet1']

print("【入力されたデータの確認】")
print("="*80)

# データを表示
for row_idx in range(5, 18):
    row_data = []
    for col_idx in range(1, 9):
        cell = ws.cell(row=row_idx, column=col_idx)
        value = cell.value if cell.value is not None else ""
        row_data.append(str(value))
    print(f"R{row_idx:02d}: " + " | ".join(row_data))

print("="*80)
print("\n✓ データ入力完了")
print(f"✓ ファイル: {excel_path}")
