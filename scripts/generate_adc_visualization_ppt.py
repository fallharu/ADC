"""ADC System の実データから、情報可視化 PowerPoint を生成する。

データソースは db/my_app_data.db（read-only）。既存の高解像度可視化から
必要なパネルだけを切り出し、主要数値・比較図は PowerPoint ネイティブ図形で
再構成する。
"""

from __future__ import annotations

import math
import random
import sqlite3
import statistics
from dataclasses import dataclass
from pathlib import Path

import cv2
import numpy as np
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "db" / "my_app_data.db"
VIS_DIR = ROOT / "output" / "database_visualization"
OUT_DIR = ROOT / "output" / "presentation"
ASSET_DIR = OUT_DIR / "assets"
OUT_PATH = OUT_DIR / "ADC_交通映像解析_情報可視化_可視化拡充版.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Yu Gothic"
FONT_NUM = "Aptos Display"


def color(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


NAVY = color("102A43")
NAVY_2 = color("163B5C")
BLUE = color("2F80ED")
BLUE_DARK = color("1D5E9F")
CYAN = color("32C5FF")
TEAL = color("1CA6A3")
GREEN = color("27AE60")
ORANGE = color("F2994A")
RED = color("E05252")
AMBER = color("F2C94C")
PURPLE = color("7656C9")
INK = color("183B56")
MID = color("627D98")
MUTED = color("829AB1")
LIGHT = color("F4F7FB")
PALE_BLUE = color("EAF3FD")
PALE_TEAL = color("E8F7F5")
PALE_ORANGE = color("FFF1E4")
PALE_RED = color("FDECEC")
WHITE = color("FFFFFF")
LINE = color("D7E1EB")
WIDENED = color("2878B5")
UNWIDENED = color("E07A2D")


@dataclass
class Stats:
    detections: int
    videos: int
    runs: int
    overtakes: int
    completed: int
    processing: int
    errors: int
    class_counts: list[tuple[str, int]]
    class_confidence: dict[str, float]
    road_videos: dict[str, int]
    road_runs: dict[str, int]
    road_detections: dict[str, int]
    road_confidence: dict[str, float]
    road_events: dict[str, int]
    road_hours: dict[str, float]
    clearance: dict[str, list[float]]
    speed_nonnull: int
    speed_zero: int
    ttc_nonnull: int
    ttc_zero: int


def load_stats() -> Stats:
    conn = sqlite3.connect(f"file:{DB_PATH.as_posix()}?mode=ro", uri=True)
    try:
        detections = conn.execute("SELECT COUNT(*) FROM DetectionRaw").fetchone()[0]
        videos = conn.execute("SELECT COUNT(*) FROM Video").fetchone()[0]
        runs = conn.execute("SELECT COUNT(*) FROM ProcessLog").fetchone()[0]
        overtakes = conn.execute("SELECT COUNT(*) FROM OvertakeEvents").fetchone()[0]
        status = dict(conn.execute("SELECT status, COUNT(*) FROM ProcessLog GROUP BY status"))
        class_rows = conn.execute(
            "SELECT class_name, COUNT(*), AVG(confidence) FROM DetectionRaw "
            "GROUP BY class_name ORDER BY COUNT(*) DESC"
        ).fetchall()

        road_videos: dict[str, int] = {}
        road_hours: dict[str, float] = {}
        for road, n, duration in conn.execute(
            "SELECT road_type, COUNT(*), SUM(duration) FROM Video "
            "WHERE road_type IS NOT NULL GROUP BY road_type"
        ):
            road_videos[road] = n
            road_hours[road] = (duration or 0) / 3600

        road_runs: dict[str, int] = {}
        road_detections: dict[str, int] = {}
        road_confidence: dict[str, float] = {}
        for road, n_runs, n_det, avg_conf in conn.execute(
            """
            SELECT v.road_type, COUNT(DISTINCT p.run_id), COUNT(d.raw_detection_id),
                   AVG(d.confidence)
            FROM Video v
            JOIN ProcessLog p ON p.video_id = v.video_id
            LEFT JOIN DetectionRaw d ON d.run_id = p.run_id
            WHERE v.road_type IS NOT NULL
            GROUP BY v.road_type
            """
        ):
            road_runs[road] = n_runs
            road_detections[road] = n_det
            road_confidence[road] = avg_conf or 0

        road_events = dict(
            conn.execute(
                """
                SELECT v.road_type, COUNT(*)
                FROM OvertakeEvents o
                JOIN ProcessLog p ON p.run_id = o.run_id
                JOIN Video v ON v.video_id = p.video_id
                WHERE v.road_type IS NOT NULL
                GROUP BY v.road_type
                """
            )
        )

        clearance: dict[str, list[float]] = {}
        for road in ("拡幅", "未拡幅"):
            clearance[road] = [
                row[0]
                for row in conn.execute(
                    """
                    SELECT d.clearance_distance_m
                    FROM Detection d
                    JOIN ProcessLog p ON p.run_id = d.run_id
                    JOIN Video v ON v.video_id = p.video_id
                    WHERE v.road_type = ?
                      AND d.clearance_distance_m BETWEEN 0.05 AND 10
                    ORDER BY d.clearance_distance_m
                    """,
                    (road,),
                )
            ]

        speed_nonnull, speed_zero, ttc_nonnull, ttc_zero = conn.execute(
            "SELECT SUM(speed_km_h IS NOT NULL), SUM(speed_km_h = 0), "
            "SUM(ttc_s IS NOT NULL), SUM(ttc_s = 0) FROM Detection"
        ).fetchone()

        return Stats(
            detections=detections,
            videos=videos,
            runs=runs,
            overtakes=overtakes,
            completed=status.get("completed", 0),
            processing=status.get("processing", 0),
            errors=status.get("error", 0),
            class_counts=[(r[0], r[1]) for r in class_rows],
            class_confidence={r[0]: r[2] for r in class_rows},
            road_videos=road_videos,
            road_runs=road_runs,
            road_detections=road_detections,
            road_confidence=road_confidence,
            road_events=road_events,
            road_hours=road_hours,
            clearance=clearance,
            speed_nonnull=speed_nonnull or 0,
            speed_zero=speed_zero or 0,
            ttc_nonnull=ttc_nonnull or 0,
            ttc_zero=ttc_zero or 0,
        )
    finally:
        conn.close()


def mann_whitney_p(x: list[float], y: list[float]) -> float:
    """連続性補正・tie補正付きの両側漸近p値。"""
    combined = sorted([(v, 0) for v in x] + [(v, 1) for v in y])
    ranks = [0.0] * len(combined)
    ties: list[int] = []
    i = 0
    while i < len(combined):
        j = i + 1
        while j < len(combined) and combined[j][0] == combined[i][0]:
            j += 1
        rank = (i + 1 + j) / 2
        for k in range(i, j):
            ranks[k] = rank
        ties.append(j - i)
        i = j
    n1, n2 = len(x), len(y)
    r1 = sum(rank for rank, item in zip(ranks, combined) if item[1] == 0)
    u1 = r1 - n1 * (n1 + 1) / 2
    n = n1 + n2
    variance = n1 * n2 / 12 * (
        (n + 1) - sum(t**3 - t for t in ties) / (n * (n - 1))
    )
    z = (abs(u1 - n1 * n2 / 2) - 0.5) / math.sqrt(variance)
    return math.erfc(z / math.sqrt(2))


def bootstrap_median_diff_ci(
    x: list[float],
    y: list[float],
    *,
    iterations: int = 20_000,
    seed: int = 20260730,
) -> tuple[float, float]:
    """中央値差（x-y）のpercentile bootstrap 95%区間。"""
    rng = random.Random(seed)
    diffs = []
    for _ in range(iterations):
        xb = [rng.choice(x) for _ in x]
        yb = [rng.choice(y) for _ in y]
        diffs.append(statistics.median(xb) - statistics.median(yb))
    diffs.sort()
    return diffs[int(iterations * 0.025)], diffs[int(iterations * 0.975) - 1]


def probability_superiority(x: list[float], y: list[float]) -> float:
    """無作為に1件ずつ選んだときx>yとなる確率（同値は0.5）。"""
    pairs = len(x) * len(y)
    greater = sum(a > b for a in x for b in y)
    ties = sum(a == b for a in x for b in y)
    return (greater + 0.5 * ties) / pairs


def wilson_interval(successes: int, total: int, z: float = 1.96) -> tuple[float, float]:
    """二項比率のWilson 95%信頼区間。"""
    p = successes / total
    denominator = 1 + z**2 / total
    center = (p + z**2 / (2 * total)) / denominator
    half = z * math.sqrt(p * (1 - p) / total + z**2 / (4 * total**2)) / denominator
    return center - half, center + half


def prepare_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    def crop(
        name: str,
        source: Path,
        box: tuple[int, int, int, int],
        contrast: float = 1.0,
        redact: tuple[tuple[int, int, int, int], ...] = (),
    ) -> Path:
        target = ASSET_DIR / name
        with Image.open(source) as image:
            item = image.convert("RGB").crop(box)
            if contrast != 1.0:
                item = ImageEnhance.Contrast(item).enhance(contrast)
            for redact_box in redact:
                blurred = item.crop(redact_box).filter(ImageFilter.GaussianBlur(radius=8))
                item.paste(blurred, redact_box)
            item.save(target, quality=94)
        return target

    heat = VIS_DIR / "bbox_detection_heatmaps.png"
    clearance = VIS_DIR / "clearance_heatmap_with_photos.png"
    overlay_road = VIS_DIR / "bbox_heatmap_on_representative_frames.png"
    overlay_class = VIS_DIR / "bbox_heatmap_car_vs_bicycle.png"
    verification = ROOT / "output" / "verification" / "run_100104"
    annotated_video = (
        ROOT / "output" / "1_250721"
        / "1_20250721_000G2603_bike_clips_20250926_183732"
        / "annotated_1_20250721_000G2603_bike_clips.mp4"
    )

    def extract_video_frame(name: str, second: float) -> Path:
        target = ASSET_DIR / name
        capture = cv2.VideoCapture(str(annotated_video))
        capture.set(cv2.CAP_PROP_POS_MSEC, second * 1000)
        ok, frame = capture.read()
        capture.release()
        if not ok:
            raise RuntimeError(f"動画フレームを抽出できません: {annotated_video} @ {second}s")
        image = Image.fromarray(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))
        image.save(target, quality=92)
        return target

    def build_full_road_heatmaps() -> dict[str, Path]:
        """拡幅・未拡幅の全DetectionRawを使った位置ヒートマップを生成する。"""
        conn = sqlite3.connect(f"file:{DB_PATH.as_posix()}?mode=ro", uri=True)
        try:
            rows = conn.execute(
                """
                WITH bounds AS (
                    SELECT run_id, MAX(x2) width, MAX(y2) height
                    FROM DetectionRaw
                    WHERE x2 > 0 AND y2 > 0
                    GROUP BY run_id
                )
                SELECT v.road_type,
                       ((d.x1+d.x2)/2.0)/b.width nx,
                       ((d.y1+d.y2)/2.0)/b.height ny
                FROM DetectionRaw d
                JOIN bounds b ON b.run_id=d.run_id
                JOIN ProcessLog p ON p.run_id=d.run_id
                JOIN Video v ON v.video_id=p.video_id
                WHERE v.road_type IN ('拡幅','未拡幅')
                  AND b.width>0 AND b.height>0
                  AND d.x2>d.x1 AND d.y2>d.y1
                """
            ).fetchall()
        finally:
            conn.close()

        groups: dict[str, list[tuple[float, float]]] = {"拡幅": [], "未拡幅": []}
        for road, nx, ny in rows:
            if nx is not None and ny is not None and 0 <= nx <= 1.05 and 0 <= ny <= 1.05:
                groups[road].append((nx, ny))

        histograms = {}
        for road, points in groups.items():
            xy = np.asarray(points, dtype=float)
            hist, _, _ = np.histogram2d(
                xy[:, 0], xy[:, 1], bins=(50, 34), range=((0, 1), (0, 1))
            )
            histograms[road] = hist / max(hist.sum(), 1) * 100
        positives = np.concatenate([h[h > 0] for h in histograms.values()])
        vmax = float(np.percentile(positives, 99.5))

        outputs = {}
        for road, hist in histograms.items():
            scaled = np.clip(hist.T / vmax, 0, 1)
            heat_u8 = (scaled * 255).astype(np.uint8)
            heat_bgr = cv2.applyColorMap(heat_u8, cv2.COLORMAP_TURBO)
            heat_rgb = cv2.cvtColor(heat_bgr, cv2.COLOR_BGR2RGB)
            heat_image = Image.fromarray(heat_rgb).resize((1400, 820), Image.Resampling.BILINEAR)
            canvas = Image.new("RGB", (1500, 930), "white")
            canvas.paste(heat_image, (55, 35))
            draw = ImageDraw.Draw(canvas)
            for i in range(1, 10):
                xx = 55 + int(1400 * i / 10)
                yy = 35 + int(820 * i / 10)
                draw.line((xx, 35, xx, 855), fill=(255, 255, 255), width=1)
                draw.line((55, yy, 1455, yy), fill=(255, 255, 255), width=1)
            draw.text((55, 872), "0%", fill=(70, 90, 110))
            draw.text((1408, 872), "100%", fill=(70, 90, 110))
            target = ASSET_DIR / ("heat_widened_full.jpg" if road == "拡幅" else "heat_unwidened_full.jpg")
            canvas.save(target, quality=93)
            outputs[road] = target
        return outputs

    full_road_heatmaps = build_full_road_heatmaps()

    return {
        # bbox_detection_heatmaps.png: 2005 x 1920
        "heat_all": crop("heat_all.jpg", heat, (120, 65, 970, 555), 1.05),
        "heat_widened": full_road_heatmaps["拡幅"],
        "heat_unwidened": full_road_heatmaps["未拡幅"],
        "heat_car": crop("heat_car.jpg", heat, (1010, 675, 1885, 1180), 1.05),
        "heat_bicycle": crop("heat_bicycle.jpg", heat, (120, 1270, 970, 1760), 1.05),
        "overlay_road": overlay_road,
        "overlay_class": overlay_class,
        "calibration_grid": verification / "frame_120_grid.jpg",
        "lane_width_graph": verification / "lane_width_graph.png",
        "scale_graph": verification / "scale_graph.png",
        "model_frame_early": extract_video_frame("model_frame_2s.jpg", 2.0),
        "model_frame_dense": extract_video_frame("model_frame_8s.jpg", 8.0),
        # clearance_heatmap_with_photos.png: 2427 x 1784。個人名行を除き、実写領域だけを抽出。
        "event_widened": crop(
            "event_widened.jpg", clearance, (319, 300, 1264, 820), 1.04,
            redact=((506, 319, 548, 344),),
        ),
        "event_unwidened": crop(
            "event_unwidened.jpg", clearance, (1469, 300, 2414, 820), 1.04,
            redact=((452, 296, 501, 326),),
        ),
    }


def add_rect(slide, x, y, w, h, fill=WHITE, radius=True, line=None, line_width=1.0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    fill=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    font=FONT,
    margin=0.0,
    italic=False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = fill
    return box


def add_multiline(slide, lines, x, y, w, h, size=16, fill=INK, bullet=False, gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if bullet else "") + line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = fill
        p.space_after = Pt(gap)
    return box


def add_title(slide, title, subtitle=None, section=None, dark=False):
    main = WHITE if dark else NAVY
    muted = color("C8D8E8") if dark else MID
    if section:
        add_text(slide, section.upper(), 0.58, 0.24, 2.1, 0.25, 9.5, CYAN if dark else BLUE, True)
    add_text(slide, title, 0.58, 0.52, 12.0, 0.54, 27, main, True, valign=MSO_ANCHOR.TOP)
    if subtitle:
        add_text(slide, subtitle, 0.6, 1.09, 12.0, 0.34, 11.5, muted, False, valign=MSO_ANCHOR.TOP)


def add_footer(slide, number, source="Source: db/my_app_data.db（read-only集計）", dark=False):
    c = color("9EB4C8") if dark else MUTED
    add_text(slide, source, 0.58, 7.17, 10.8, 0.18, 8.2, c, False, valign=MSO_ANCHOR.BOTTOM)
    add_text(slide, f"{number:02d}", 12.25, 7.12, 0.48, 0.22, 9, c, True, PP_ALIGN.RIGHT)


def add_pill(slide, text, x, y, w, fill, text_color=WHITE, size=10.5):
    add_rect(slide, x, y, w, 0.34, fill=fill, radius=True)
    add_text(slide, text, x, y + 0.01, w, 0.30, size, text_color, True, PP_ALIGN.CENTER)


def add_kpi(slide, x, y, w, h, value, label, accent=BLUE, note=None):
    add_rect(slide, x + 0.04, y + 0.06, w, h, fill=LINE, radius=True)
    add_rect(slide, x, y, w, h, fill=WHITE, radius=True)
    add_rect(slide, x, y, 0.08, h, fill=accent, radius=False)
    add_text(slide, value, x + 0.22, y + 0.16, w - 0.35, 0.55, 26, NAVY, True, font=FONT_NUM)
    add_text(slide, label, x + 0.22, y + 0.72, w - 0.35, 0.26, 11, MID, True)
    if note:
        add_text(slide, note, x + 0.22, y + 1.00, w - 0.35, 0.24, 8.5, MUTED)


def add_picture(slide, path: Path, x, y, w, h, line=LINE):
    add_rect(slide, x + 0.04, y + 0.05, w, h, fill=LINE, radius=True)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if line:
        pic.line.color.rgb = line
        pic.line.width = Pt(0.8)
    return pic


def add_line(slide, x1, y1, x2, y2, line_color=LINE, width=1.0):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x1),
        Inches(y1),
        Inches(max(x2 - x1, 0.01)),
        Inches(max(y2 - y1, 0.01)),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = line_color
    line.line.fill.background()
    return line


def add_circle(slide, x, y, d, fill, text=None, text_color=WHITE, size=14, line=None):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    circle.fill.solid()
    circle.fill.fore_color.rgb = fill
    if line:
        circle.line.color.rgb = line
    else:
        circle.line.fill.background()
    if text is not None:
        tf = circle.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(text)
        run.font.name = FONT_NUM
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = text_color
    return circle


def add_stacked_bar(slide, x, y, w, h, parts, labels, colors, label_left, label_color):
    add_text(slide, label_left, x, y - 0.03, 1.1, h, 12.5, label_color, True)
    bar_x = x + 1.2
    total = sum(parts)
    cursor = bar_x
    for value, label, fill in zip(parts, labels, colors):
        width = (w - 1.2) * value / total
        add_rect(slide, cursor, y, width, h, fill=fill, radius=False)
        if width > 0.65:
            add_text(slide, f"{value / total * 100:.1f}%", cursor, y, width, h, 10, WHITE if fill != AMBER else NAVY, True, PP_ALIGN.CENTER)
        cursor += width


def make_presentation(stats: Stats, assets: dict[str, Path]) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "単眼カメラ×AIで、追越し安全性を可視化"
    prs.core_properties.subject = "ADC System 実データ情報可視化"
    prs.core_properties.author = "ADC System"
    prs.core_properties.keywords = "ADC, 交通安全, 情報可視化, 追越し, 離隔距離"
    prs.core_properties.comments = "db/my_app_data.db の read-only 集計を使用。2026-07-30可視化拡充版。"
    blank = prs.slide_layouts[6]

    widened = stats.clearance["拡幅"]
    unwidened = stats.clearance["未拡幅"]
    p_value = mann_whitney_p(widened, unwidened)
    widened_mean = statistics.mean(widened)
    unwidened_mean = statistics.mean(unwidened)
    widened_median = statistics.median(widened)
    unwidened_median = statistics.median(unwidened)
    widened_low = sum(v < 1.5 for v in widened) / len(widened)
    unwidened_low = sum(v < 1.5 for v in unwidened) / len(unwidened)
    widened_low_n = sum(v < 1.5 for v in widened)
    unwidened_low_n = sum(v < 1.5 for v in unwidened)
    median_diff_ci = bootstrap_median_diff_ci(widened, unwidened)
    superiority = probability_superiority(widened, unwidened)
    widened_low_ci = wilson_interval(widened_low_n, len(widened))
    unwidened_low_ci = wilson_interval(unwidened_low_n, len(unwidened))
    widened_rate = stats.road_events["拡幅"] / stats.road_hours["拡幅"]
    unwidened_rate = stats.road_events["未拡幅"] / stats.road_hours["未拡幅"]
    class_count_map = dict(stats.class_counts)
    motor_vehicle_records = sum(class_count_map.get(name, 0) for name in ("car", "truck", "bus"))
    tire_records = class_count_map.get("Tire", 0)
    bicycle_records = class_count_map.get("bicycle", 0)
    bicycle_tire_records = class_count_map.get("Bicycle_Tires", 0)

    # 01 — Cover
    slide = prs.slides.add_slide(blank)
    bg = add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY, radius=False)
    add_rect(slide, 8.16, 0, 5.17, 7.5, fill=NAVY_2, radius=False)
    # データの流れを示す装飾線
    for i, c in enumerate((CYAN, BLUE, TEAL, ORANGE)):
        add_rect(slide, 7.62 + i * 0.15, 0, 0.055, 7.5, fill=c, radius=False)
    add_picture(slide, assets["heat_all"], 8.45, 0.72, 4.35, 2.52, line=None)
    add_picture(slide, assets["heat_car"], 8.45, 3.55, 4.35, 2.52, line=None)
    add_pill(slide, "ADC SYSTEM / DATA STORY", 0.72, 0.72, 2.55, BLUE)
    add_text(slide, "単眼カメラ × AIで、", 0.72, 1.52, 6.55, 0.62, 29, color("D8E9F8"), True, valign=MSO_ANCHOR.TOP)
    add_text(slide, "追越し安全性を\n可視化する", 0.72, 2.08, 6.60, 1.54, 40, WHITE, True, valign=MSO_ANCHOR.TOP)
    add_text(slide, "交通映像を「件数」から「安全余裕」のデータへ", 0.76, 3.84, 6.48, 0.45, 17, CYAN, True)
    add_text(slide, "169動画・121万検出レコードを用いた\nADC交通映像解析システムの実データ可視化", 0.76, 4.48, 6.5, 0.78, 14.5, color("C8D8E8"), False, valign=MSO_ANCHOR.TOP)
    add_line(slide, 0.76, 6.55, 6.9, 6.58, color("355777"))
    add_text(slide, "2026.07  |  SQLite analysis snapshot", 0.76, 6.66, 6.2, 0.25, 10, color("9EB4C8"), True)

    # 02 — Executive summary
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=LIGHT, radius=False)
    add_title(slide, "可視化で、検出位置・クラス構造・安全余裕が読める", "結果だけでなく処理の途中も見せることで、『どこに・何が・どの余裕でいたか』を説明可能に", "WHAT BECAME VISIBLE")
    cards = [
        (0.58, "全検出の位置分布", "拡幅  vs  未拡幅", "BBox中心を画面内で正規化", "道路空間を比較", BLUE,
         stats.road_detections["拡幅"] / max(stats.road_detections.values()), stats.road_detections["未拡幅"] / max(stats.road_detections.values())),
        (4.48, "クラス構造", f"Tire {tire_records/1000:.0f}k", f"car {class_count_map['car']/1000:.0f}k｜大型車も対象", "想定どおり", PURPLE,
         tire_records / max(tire_records, class_count_map["car"]), class_count_map["car"] / max(tire_records, class_count_map["car"])),
        (8.38, "離隔距離の中央値", f"{widened_median:.2f} m  vs  {unwidened_median:.2f} m", "拡幅 / 未拡幅", f"差 +{widened_median-unwidened_median:.2f}m", ORANGE,
         widened_median / max(widened_median, unwidened_median), unwidened_median / max(widened_median, unwidened_median)),
    ]
    for x, label, value, unit, takeaway, accent, ratio_w, ratio_u in cards:
        mini_a, mini_b = (("Tire", "car") if x == 4.48 else ("拡幅", "未拡幅"))
        add_rect(slide, x + 0.05, 1.66, 3.68, 3.65, fill=LINE, radius=True)
        add_rect(slide, x, 1.60, 3.68, 3.65, fill=WHITE, radius=True)
        add_rect(slide, x, 1.60, 3.68, 0.11, fill=accent, radius=False)
        add_text(slide, label, x + 0.24, 1.92, 3.18, 0.36, 14, MID, True)
        add_text(slide, value, x + 0.24, 2.42, 3.18, 0.70, 25, NAVY, True, PP_ALIGN.CENTER, font=FONT_NUM)
        add_text(slide, unit, x + 0.24, 3.12, 3.18, 0.30, 10.5, MUTED, False, PP_ALIGN.CENTER)
        add_pill(slide, takeaway, x + 0.65, 3.74, 2.38, accent, size=11.5)
        # 比較のミニバー
        add_text(slide, mini_a, x + 0.25, 4.42, 0.55, 0.20, 8.5, WIDENED, True)
        add_rect(slide, x + 0.87, 4.46, 2.35, 0.12, fill=PALE_BLUE, radius=False)
        add_rect(slide, x + 0.87, 4.46, 2.35 * ratio_w, 0.12, fill=WIDENED, radius=False)
        add_text(slide, mini_b, x + 0.25, 4.76, 0.55, 0.20, 8.5, UNWIDENED, True)
        add_rect(slide, x + 0.87, 4.80, 2.35, 0.12, fill=PALE_ORANGE, radius=False)
        add_rect(slide, x + 0.87, 4.80, 2.35 * ratio_u, 0.12, fill=UNWIDENED, radius=False)
    add_rect(slide, 0.58, 5.72, 11.48, 0.78, fill=PALE_BLUE, radius=True)
    add_circle(slide, 0.82, 5.92, 0.34, BLUE, "!", WHITE, 13)
    add_text(slide, "読み方：可視化は『正しさの証明』ではなく、検出・校正・集計の状態と、比較で見える傾向を確認するための共通画面。", 1.28, 5.84, 10.35, 0.48, 13.0, NAVY, True)
    add_footer(slide, 2, "Source: 現行DB｜全検出・クラス別レコード・有効離隔43件")

    # 03 — Problem framing
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=WHITE, radius=False)
    add_title(slide, "件数だけでは、事故前の“安全余裕”が見えない", "従来の目視調査を、距離・軌跡・追越し文脈まで扱うデータ基盤へ", "WHY ADC")
    add_rect(slide, 0.62, 1.58, 4.66, 4.74, fill=PALE_ORANGE, radius=True)
    add_pill(slide, "従来：目視・手作業", 0.96, 1.90, 1.90, ORANGE)
    add_text(slide, "“何台通ったか”", 0.96, 2.52, 3.90, 0.54, 28, NAVY, True)
    add_multiline(slide, ["件数・車種が中心", "観測者間で判断がばらつく", "微小な距離変化を追いにくい", "大量動画の再確認に時間"], 0.98, 3.28, 3.9, 1.75, 15, INK, bullet=True, gap=8)
    add_rect(slide, 1.02, 5.35, 3.78, 0.55, fill=WHITE, radius=True)
    add_text(slide, "結果：危険の“前兆”が残りにくい", 1.15, 5.42, 3.50, 0.34, 12.5, RED, True, PP_ALIGN.CENTER)

    add_circle(slide, 5.58, 3.42, 0.76, BLUE, "→", WHITE, 23)

    add_rect(slide, 6.58, 1.58, 6.12, 4.74, fill=PALE_BLUE, radius=True)
    add_pill(slide, "ADC：AI＋人の検証", 6.94, 1.90, 2.10, BLUE)
    add_text(slide, "“どれだけ余裕があったか”", 6.94, 2.52, 5.22, 0.54, 28, NAVY, True)
    for i, (title, note, accent) in enumerate([
        ("検出・追跡", "車・自転車をフレーム横断でID化", BLUE),
        ("幾何計測", "白線・接近・離隔を実距離へ変換", TEAL),
        ("追越し文脈", "誰が誰を、いつ、どの余裕で追越したか", ORANGE),
    ]):
        yy = 3.25 + i * 0.80
        add_circle(slide, 6.98, yy, 0.42, accent, str(i + 1), WHITE, 11)
        add_text(slide, title, 7.55, yy - 0.03, 1.36, 0.32, 13, NAVY, True)
        add_text(slide, note, 8.96, yy - 0.03, 3.20, 0.38, 11.2, MID)
    add_rect(slide, 6.94, 5.65, 5.28, 0.45, fill=WHITE, radius=True)
    add_text(slide, "目的：件数調査 → 道路改善の定量根拠", 7.10, 5.69, 4.96, 0.30, 13.2, BLUE_DARK, True, PP_ALIGN.CENTER)
    add_footer(slide, 3, "Source: docs/ADC_System_Overview.md / 実装構成を要約")

    # 04 — Workflow
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=LIGHT, radius=False)
    add_title(slide, "動画から安全指標までを、一気通貫で処理", "キャリブレーションとHuman-in-the-loopが、映像解析を監査可能なデータへ変える", "SYSTEM FLOW")
    steps = [
        ("01", "動画入力", "MP4 / MOV\nフォルダ一括", BLUE),
        ("02", "幾何校正", "白線・距離\nカメラ条件", PURPLE),
        ("03", "検出・追跡", "YOLO\nByteTrack", TEAL),
        ("04", "指標算出", "離隔・接近\n白線・追越し", ORANGE),
        ("05", "人の検証", "追加・削除\n修正履歴", RED),
        ("06", "出力", "Excel / CSV\n注釈動画", GREEN),
    ]
    sx, sy, sw, gap = 0.62, 1.85, 1.76, 0.25
    for i, (num, name, note, accent) in enumerate(steps):
        x = sx + i * (sw + gap)
        add_rect(slide, x + 0.04, sy + 0.05, sw, 3.05, fill=LINE, radius=True)
        add_rect(slide, x, sy, sw, 3.05, fill=WHITE, radius=True)
        add_circle(slide, x + 0.53, sy + 0.28, 0.70, accent, num, WHITE, 14)
        add_text(slide, name, x + 0.15, sy + 1.18, sw - 0.30, 0.42, 15, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, note, x + 0.16, sy + 1.78, sw - 0.32, 0.74, 11.2, MID, False, PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
        if i < len(steps) - 1:
            add_text(slide, "›", x + sw + 0.03, sy + 1.20, gap - 0.02, 0.50, 27, MUTED, True, PP_ALIGN.CENTER)
    add_rect(slide, 0.62, 5.40, 12.08, 0.88, fill=NAVY, radius=True)
    add_text(slide, "AIが候補を広く拾い、人がイベントを確定。修正後の指標を再計算し、根拠と履歴を残す。", 0.98, 5.55, 11.40, 0.42, 15, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide, 4, "Source: Source_code/modules・routes の実装構成を要約")

    # 05 — Model output visualization
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=LIGHT, radius=False)
    add_title(slide, "モデルが『何を見て、どう追っているか』をフレーム上で確認", "BBox・追跡ID・クラス・距離情報を重ね、数値の元になった認識状態を人が読める形にする", "MODEL VIEW")
    add_rect(slide, 0.62, 1.50, 5.88, 3.82, fill=WHITE, radius=True)
    add_rect(slide, 6.82, 1.50, 5.88, 3.82, fill=WHITE, radius=True)
    add_picture(slide, assets["model_frame_early"], 0.84, 1.95, 5.44, 3.06)
    add_picture(slide, assets["model_frame_dense"], 7.04, 1.95, 5.44, 3.06)
    add_pill(slide, "少数対象｜追跡線を確認", 0.96, 1.66, 1.92, BLUE, size=10)
    add_pill(slide, "複数対象｜重なりを確認", 7.16, 1.66, 2.05, ORANGE, size=10)

    layers = [
        ("BBox", "対象の位置・大きさ", BLUE),
        ("Track ID", "フレームをまたぐ同一性", PURPLE),
        ("Class", "car / truck / tire 等", TEAL),
        ("Distance", "接近・離隔の計測値", ORANGE),
    ]
    for i, (name, note, accent) in enumerate(layers):
        x = 0.74 + i * 3.02
        add_rect(slide, x, 5.62, 2.76, 0.64, fill=WHITE, radius=True)
        add_circle(slide, x + 0.16, 5.77, 0.31, accent, str(i + 1), WHITE, 8.5)
        add_text(slide, name, x + 0.58, 5.68, 0.83, 0.24, 10.5, NAVY, True)
        add_text(slide, note, x + 1.38, 5.68, 1.18, 0.30, 8.7, MID, False, PP_ALIGN.CENTER)
    add_footer(slide, 5, "Source: 実注釈動画から代表フレームを抽出｜検出・追跡・距離表示")

    # 06 — Calibration visualization
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=WHITE, radius=False)
    add_title(slide, "画面上のピクセルを、道路上の距離へ変換する過程も可視化", "道路グリッド・左右境界・奥行き別スケールを重ね、計測の前提をレビューできる", "GEOMETRY VIEW")
    add_rect(slide, 0.62, 1.48, 7.35, 4.98, fill=LIGHT, radius=True)
    add_picture(slide, assets["calibration_grid"], 0.83, 1.83, 6.93, 3.90)
    add_pill(slide, "道路グリッド＋境界線", 1.00, 1.65, 1.80, TEAL, size=10)
    add_text(slide, "黄色点：距離グリッド　青・緑：道路境界　赤：基準線", 1.02, 5.85, 6.55, 0.30, 10.5, MID, True, PP_ALIGN.CENTER)

    add_rect(slide, 8.25, 1.48, 4.45, 2.30, fill=LIGHT, radius=True)
    add_picture(slide, assets["lane_width_graph"], 8.50, 1.82, 3.95, 1.64)
    add_pill(slide, "画面奥行き × 道路幅", 8.62, 1.61, 1.74, BLUE, size=9.5)
    add_rect(slide, 8.25, 4.05, 4.45, 2.41, fill=LIGHT, radius=True)
    add_picture(slide, assets["scale_graph"], 8.50, 4.40, 3.95, 1.72)
    add_pill(slide, "画面奥行き × 1m尺度", 8.62, 4.18, 1.78, PURPLE, size=9.5)
    add_footer(slide, 6, "Source: output/verification/run_100104｜幾何校正の検証出力")

    # 07 — Data foundation
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=LIGHT, radius=False)
    add_title(slide, "169動画から、121万件の検出レコードを蓄積", "検出レコードはユニーク物体数ではなく、フレーム単位の観測データ", "DATA FOUNDATION")
    add_kpi(slide, 0.62, 1.55, 2.84, 1.34, f"{stats.videos:,}", "動画", BLUE)
    add_kpi(slide, 3.68, 1.55, 2.84, 1.34, f"{stats.detections/1_000_000:.2f}M", "検出レコード", TEAL, f"{stats.detections:,}件")
    add_kpi(slide, 6.74, 1.55, 2.84, 1.34, f"{stats.runs:,}", "処理Run", PURPLE)
    add_kpi(slide, 9.80, 1.55, 2.84, 1.34, f"{stats.overtakes:,}", "追越しイベント", ORANGE, "有効離隔 43件")

    add_rect(slide, 0.62, 3.20, 7.15, 3.34, fill=WHITE, radius=True)
    add_text(slide, "クラス別レコード数", 0.94, 3.45, 3.0, 0.32, 14, NAVY, True)
    max_count = max(v for _, v in stats.class_counts)
    class_colors = {"Tire": PURPLE, "car": BLUE, "truck": ORANGE, "Bicycle_Tires": color("9B7BD6"), "bicycle": TEAL, "bus": GREEN}
    label_map = {"Tire": "Tire", "car": "car", "truck": "truck", "Bicycle_Tires": "Bicycle_Tires", "bicycle": "bicycle", "bus": "bus"}
    for i, (name, value) in enumerate(stats.class_counts):
        yy = 3.93 + i * 0.39
        add_text(slide, label_map.get(name, name), 0.94, yy, 1.10, 0.22, 9.2, MID, True)
        add_rect(slide, 2.08, yy + 0.045, 4.22, 0.14, fill=color("EDF2F7"), radius=False)
        add_rect(slide, 2.08, yy + 0.045, 4.22 * value / max_count, 0.14, fill=class_colors.get(name, BLUE), radius=False)
        add_text(slide, f"{value:,}", 6.38, yy - 0.02, 0.95, 0.24, 9.2, INK, True, PP_ALIGN.RIGHT, font=FONT_NUM)

    add_rect(slide, 8.04, 3.20, 4.66, 3.34, fill=WHITE, radius=True)
    completion = stats.completed / stats.runs
    add_text(slide, "処理状態", 8.37, 3.45, 1.8, 0.32, 14, NAVY, True)
    add_text(slide, f"{completion*100:.1f}%", 8.37, 3.89, 2.02, 0.62, 29, GREEN, True, font=FONT_NUM)
    add_text(slide, "完了", 10.25, 4.05, 0.78, 0.26, 11, MID, True)
    add_rect(slide, 8.37, 4.70, 3.84, 0.22, fill=color("E8EEF4"), radius=True)
    add_rect(slide, 8.37, 4.70, 3.84 * completion, 0.22, fill=GREEN, radius=True)
    for i, (label, value, accent) in enumerate([
        ("completed", stats.completed, GREEN),
        ("processing", stats.processing, AMBER),
        ("error", stats.errors, RED),
    ]):
        yy = 5.22 + i * 0.35
        add_circle(slide, 8.38, yy + 0.04, 0.16, accent)
        add_text(slide, label, 8.66, yy, 1.40, 0.23, 9.5, MID)
        add_text(slide, str(value), 11.42, yy, 0.58, 0.23, 10, NAVY, True, PP_ALIGN.RIGHT, font=FONT_NUM)
    add_footer(slide, 7)

    # 08 — Class record structure
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=LIGHT, radius=False)
    add_title(slide, "タイヤ検出が車体より多いのは、モデル構造として想定どおり", "1台の車両に複数のタイヤがあり、Tireにはcarだけでなくtruck・busのタイヤも含まれる", "CLASS STRUCTURE")
    class_panels = [
        (
            0.62, "自動車系の全体", "車体レコード", motor_vehicle_records, "Tire", tire_records,
            f"Tire / (car+truck+bus) = {tire_records/motor_vehicle_records:.2f}", BLUE,
            "車体3クラスを合算しても、Tireがやや多い",
        ),
        (
            4.48, "carとの単純比較", "car", class_count_map["car"], "Tire", tire_records,
            f"Tire / car = {tire_records/class_count_map['car']:.2f}", PURPLE,
            "Tire側にはtruck・bus由来も含まれる",
        ),
        (
            8.34, "自転車系", "bicycle", bicycle_records, "Bicycle_Tires", bicycle_tire_records,
            f"Bicycle_Tires / bicycle = {bicycle_tire_records/bicycle_records:.2f}", TEAL,
            "自転車は2輪を検出できている場面が多い",
        ),
    ]
    for x, title, label_a, value_a, label_b, value_b, ratio_text, accent, reading in class_panels:
        add_rect(slide, x + 0.05, 1.68, 3.68, 3.82, fill=LINE, radius=True)
        add_rect(slide, x, 1.62, 3.68, 3.82, fill=WHITE, radius=True)
        add_rect(slide, x, 1.62, 3.68, 0.11, fill=accent, radius=False)
        add_text(slide, title, x + 0.22, 1.91, 3.24, 0.34, 14, NAVY, True, PP_ALIGN.CENTER)
        maximum = max(value_a, value_b)
        for row, (label, value, fill) in enumerate(((label_a, value_a, WIDENED), (label_b, value_b, accent))):
            yy = 2.53 + row * 0.86
            add_text(slide, label, x + 0.25, yy, 1.12, 0.24, 10.0, MID, True)
            add_text(slide, f"{value:,}", x + 1.68, yy - 0.05, 1.58, 0.34, 16, NAVY, True, PP_ALIGN.RIGHT, font=FONT_NUM)
            add_rect(slide, x + 0.25, yy + 0.39, 3.05, 0.15, fill=color("E8EEF4"), radius=False)
            add_rect(slide, x + 0.25, yy + 0.39, 3.05 * value / maximum, 0.15, fill=fill, radius=False)
        add_pill(slide, ratio_text, x + 0.58, 4.35, 2.52, accent, size=10.5)
        add_text(slide, reading, x + 0.30, 4.82, 3.08, 0.42, 9.5, MID, True, PP_ALIGN.CENTER)

    add_rect(slide, 0.62, 5.82, 12.08, 0.60, fill=NAVY, radius=True)
    add_text(slide, "重要：ここで数えているのはフレーム単位のレコード。比率は『1台あたりの正確なタイヤ数』ではなく、検出構造が想定どおりかを見る指標。", 0.94, 5.94, 11.44, 0.34, 12.0, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide, 8, "Source: DetectionRaw クラス別レコード｜Tireにはcar・truck・bus由来を含む")

    # 09 — Analysis design
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=WHITE, radius=False)
    add_title(slide, "比較の単位・分母・除外条件を明示", "121万検出レコードをそのまま標本数とせず、追越しイベント単位で記述的に比較", "ANALYSIS DESIGN")
    design_cards = [
        (
            0.62, 1.56, "01｜分析単位", BLUE,
            "Video → Run → Event",
            [f"動画 {stats.videos} / Run {stats.runs}", f"追越し {stats.overtakes}件", f"有効離隔 {len(widened)+len(unwidened)}件"],
            "フレーム観測を独立標本として扱わない",
        ),
        (
            6.79, 1.56, "02｜組入れ条件", TEAL,
            "道路区分＋有効距離",
            ["road_type が拡幅 / 未拡幅", "離隔 0.05–10m", f"拡幅 {len(widened)} / 未拡幅 {len(unwidened)}"],
            f"離隔欠測 {stats.overtakes-len(widened)-len(unwidened)}件は別管理",
        ),
        (
            0.62, 4.02, "03｜指標と分母", ORANGE,
            "頻度と余裕を分ける",
            ["頻度＝追越し件数 / 動画時間", "余裕＝イベント別の離隔距離", "中央値＋距離帯の構成比"],
            "1.5mは比較用の参考帯で、法的判定ではない",
        ),
        (
            6.79, 4.02, "04｜推論の範囲", PURPLE,
            "差の兆候を評価",
            ["Mann–Whitney U（両側）", "中央値差のbootstrap区間", "二項比率のWilson区間"],
            "地点・方向・時間・天候を未調整のため因果は断定しない",
        ),
    ]
    for x, y, label, accent, headline, lines, caution in design_cards:
        add_rect(slide, x + 0.04, y + 0.05, 5.55, 2.05, fill=LINE, radius=True)
        add_rect(slide, x, y, 5.55, 2.05, fill=LIGHT, radius=True)
        add_rect(slide, x, y, 0.10, 2.05, fill=accent, radius=False)
        add_text(slide, label, x + 0.26, y + 0.17, 1.62, 0.28, 10.5, accent, True)
        add_text(slide, headline, x + 1.88, y + 0.14, 3.36, 0.34, 15, NAVY, True, PP_ALIGN.RIGHT)
        add_multiline(slide, lines, x + 0.28, y + 0.62, 4.98, 0.82, 11.2, INK, bullet=True, gap=3)
        add_rect(slide, x + 0.26, y + 1.54, 5.03, 0.32, fill=WHITE, radius=True)
        add_text(slide, caution, x + 0.38, y + 1.57, 4.78, 0.24, 9.2, MID, True, PP_ALIGN.CENTER)
    add_footer(slide, 9, "Source: 現行DBの抽出条件・比較ロジック｜再現コード: scripts/generate_adc_visualization_ppt.py")

    # 10 — Road-type spatial comparison
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=WHITE, radius=False)
    add_title(slide, "全検出を分けると、拡幅・未拡幅で検出帯の形が異なる", "車種を限定せず、Tire系を含む全BBox中心を道路区分ごとに正規化して比較", "ROAD-TYPE FOOTPRINT")
    road_panels = [
        (0.62, assets["heat_widened"], f"拡幅｜{stats.road_detections['拡幅']:,}件", "奥行き方向へ連続する検出帯", WIDENED),
        (6.82, assets["heat_unwidened"], f"未拡幅｜{stats.road_detections['未拡幅']:,}件", "画面中央～下部に集まる検出帯", UNWIDENED),
    ]
    for x, path, label, reading, accent in road_panels:
        add_rect(slide, x, 1.52, 5.88, 3.92, fill=LIGHT, radius=True)
        add_picture(slide, path, x + 0.24, 1.92, 5.40, 3.12)
        add_pill(slide, label, x + 0.40, 1.68, 2.02, accent, size=10)
        add_text(slide, reading, x + 0.42, 5.08, 5.04, 0.26, 11.5, NAVY, True, PP_ALIGN.CENTER)
    add_rect(slide, 0.62, 5.78, 12.08, 0.64, fill=PALE_BLUE, radius=True)
    add_text(slide, "読み取れること：検出位置の違いを画像感覚ではなく分布として比較できる。一方、画角・道路線形・交通構成の影響も含むため、道路幅だけの効果とは断定しない。", 0.92, 5.88, 11.48, 0.42, 11.5, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, 10, "Source: 拡幅・未拡幅の全DetectionRaw｜全クラス｜各群内構成比")

    # 11 — Class spatial footprints
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=WHITE, radius=False)
    add_title(slide, "検出位置は、道路の奥行きと車種ごとに異なる", "BBox中心の正規化ヒートマップ。色は各パネル内の構成比で、絶対件数ではない", "SPATIAL FOOTPRINT")
    panel_info = [
        (0.62, assets["heat_all"], "全検出", "道路形状に沿う主検出帯"),
        (4.46, assets["heat_car"], "car", "画面右奥～中央に集中"),
        (8.30, assets["heat_bicycle"], "bicycle", "路肩側・画面下部に偏る"),
    ]
    for x, path, label, note in panel_info:
        add_picture(slide, path, x, 1.58, 3.55, 2.05)
        add_pill(slide, label, x + 0.18, 1.76, 0.98, NAVY, size=10)
        add_rect(slide, x, 3.88, 3.55, 1.30, fill=LIGHT, radius=True)
        add_text(slide, note, x + 0.20, 4.04, 3.15, 0.36, 13, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, "検出密度の“形”が、画角・走行帯・対象物の違いを反映", x + 0.25, 4.48, 3.05, 0.42, 9.3, MID, False, PP_ALIGN.CENTER)
    add_rect(slide, 0.62, 5.68, 11.23, 0.62, fill=PALE_TEAL, radius=True)
    add_text(slide, "示唆：道路区分比較では、検出件数だけでなく画角・方向・対象クラスを揃えて評価する必要がある。", 0.92, 5.80, 10.62, 0.35, 13, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, 11, "Source: bbox_detection_heatmaps.png｜DetectionRaw 1/5抽出・Run内正規化")

    # 12 — Heatmap on representative frames
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=LIGHT, radius=False)
    add_title(slide, "ヒートマップを元映像へ戻すと、道路上の意味が読める", "正規化分布だけでなく代表フレームへ重ね、検出帯が車線・路肩・奥行きのどこに対応するか確認", "CONTEXTUAL VIEW")
    add_rect(slide, 0.62, 1.48, 12.08, 4.66, fill=WHITE, radius=True)
    add_picture(slide, assets["overlay_road"], 0.91, 1.77, 11.50, 3.72)
    add_pill(slide, "拡幅", 1.14, 1.93, 0.90, WIDENED, size=10)
    add_pill(slide, "未拡幅", 6.95, 1.93, 1.02, UNWIDENED, size=10)
    add_text(slide, "分布図 → 元フレーム → 道路上の位置、の順に往復できるため、集計値の解釈とモデルの妥当性確認を同じ資料上で行える。", 1.04, 5.60, 11.18, 0.34, 12.0, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, 12, "Source: bbox_heatmap_on_representative_frames.png｜代表Runの全BBox中心密度")

    # 13 — Event-level metric
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=LIGHT, radius=False)
    add_title(slide, "追越しの“瞬間”を、離隔距離として残す", "代表写真は各道路区分の中央値に近い実イベント。赤点と青点の横方向距離を計測", "EVENT-LEVEL VISUALIZATION")
    add_rect(slide, 0.62, 1.48, 5.88, 3.66, fill=WHITE, radius=True)
    add_rect(slide, 6.82, 1.48, 5.88, 3.66, fill=WHITE, radius=True)
    add_picture(slide, assets["event_widened"], 0.87, 1.96, 5.38, 2.96)
    add_picture(slide, assets["event_unwidened"], 7.07, 1.96, 5.38, 2.96)
    add_pill(slide, "拡幅｜代表例 2.01 m", 0.98, 1.64, 1.92, WIDENED, size=10.5)
    add_pill(slide, "未拡幅｜代表例 1.82 m", 7.18, 1.64, 2.10, UNWIDENED, size=10.5)
    add_rect(slide, 0.62, 5.48, 12.08, 0.95, fill=NAVY, radius=True)
    flow = [("車体BBox", BLUE), ("計測点", CYAN), ("横方向の最短距離", ORANGE), ("道路区分別に集計", GREEN)]
    fx = 1.06
    for i, (label, accent) in enumerate(flow):
        add_circle(slide, fx, 5.72, 0.38, accent, str(i + 1), WHITE, 10)
        add_text(slide, label, fx + 0.49, 5.66, 1.85, 0.48, 11, WHITE, True)
        if i < len(flow) - 1:
            add_text(slide, "→", fx + 2.02, 5.66, 0.38, 0.45, 16, CYAN, True, PP_ALIGN.CENTER)
        fx += 2.85
    add_footer(slide, 13, "Source: clearance_heatmap_with_photos.png｜車両番号はぼかし済み・比較色帯は法的判定ではない")

    # 14 — Road comparison
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=WHITE, radius=False)
    add_title(slide, "未拡幅では、狭い離隔の割合が高い", "1.5mは比較用の参考帯。法的閾値や安全判定としては扱わない", "ROAD COMPARISON")
    add_rect(slide, 0.62, 1.54, 7.35, 3.35, fill=LIGHT, radius=True)
    add_text(slide, "離隔距離帯の構成比", 0.92, 1.80, 3.0, 0.32, 15, NAVY, True)
    bins = [
        sum(v < 1.5 for v in widened),
        sum(1.5 <= v < 2.0 for v in widened),
        sum(2.0 <= v < 2.5 for v in widened),
        sum(v >= 2.5 for v in widened),
    ]
    bins_u = [
        sum(v < 1.5 for v in unwidened),
        sum(1.5 <= v < 2.0 for v in unwidened),
        sum(2.0 <= v < 2.5 for v in unwidened),
        sum(v >= 2.5 for v in unwidened),
    ]
    band_labels = ["<1.5", "1.5–2.0", "2.0–2.5", "≥2.5"]
    band_colors = [RED, ORANGE, TEAL, BLUE]
    add_stacked_bar(slide, 0.96, 2.48, 6.50, 0.50, bins, band_labels, band_colors, "拡幅", WIDENED)
    add_stacked_bar(slide, 0.96, 3.46, 6.50, 0.50, bins_u, band_labels, band_colors, "未拡幅", UNWIDENED)
    lx = 1.18
    for label, accent in zip(band_labels, band_colors):
        add_circle(slide, lx, 4.36, 0.16, accent)
        add_text(slide, f"{label} m", lx + 0.24, 4.30, 1.10, 0.25, 9.2, MID)
        lx += 1.42

    add_rect(slide, 8.24, 1.54, 4.46, 3.35, fill=LIGHT, radius=True)
    add_text(slide, "比較サマリー", 8.57, 1.80, 2.3, 0.32, 15, NAVY, True)
    summaries = [
        ("中央値", f"{widened_median:.2f}m", f"{unwidened_median:.2f}m", TEAL),
        ("平均", f"{widened_mean:.2f}m", f"{unwidened_mean:.2f}m", BLUE),
        ("1.5m未満", f"{widened_low*100:.1f}%", f"{unwidened_low*100:.1f}%", RED),
    ]
    add_text(slide, "指標", 8.58, 2.26, 1.25, 0.25, 9.5, MUTED, True)
    add_text(slide, "拡幅", 10.15, 2.26, 0.76, 0.25, 9.5, WIDENED, True, PP_ALIGN.CENTER)
    add_text(slide, "未拡幅", 11.32, 2.26, 0.86, 0.25, 9.5, UNWIDENED, True, PP_ALIGN.CENTER)
    for i, (label, a, b, accent) in enumerate(summaries):
        yy = 2.70 + i * 0.57
        add_circle(slide, 8.58, yy + 0.06, 0.17, accent)
        add_text(slide, label, 8.88, yy, 1.05, 0.29, 10.5, INK, True)
        add_text(slide, a, 10.06, yy, 1.02, 0.29, 12, WIDENED, True, PP_ALIGN.CENTER, font=FONT_NUM)
        add_text(slide, b, 11.25, yy, 1.02, 0.29, 12, UNWIDENED, True, PP_ALIGN.CENTER, font=FONT_NUM)
    add_text(slide, f"Mann–Whitney U｜p = {p_value:.3f}", 8.58, 4.45, 3.55, 0.25, 9.8, MID, True, PP_ALIGN.CENTER)

    add_rect(slide, 0.62, 5.23, 12.08, 1.12, fill=NAVY, radius=True)
    add_text(slide, "追越し頻度（件 / 動画時間）", 0.94, 5.43, 2.48, 0.27, 11.5, color("C8D8E8"), True)
    add_text(slide, f"拡幅  {widened_rate:.1f}", 3.58, 5.33, 1.80, 0.47, 19, WHITE, True, PP_ALIGN.CENTER, font=FONT_NUM)
    add_text(slide, "vs", 5.47, 5.38, 0.50, 0.36, 11, CYAN, True, PP_ALIGN.CENTER)
    add_text(slide, f"未拡幅  {unwidened_rate:.1f}", 6.06, 5.33, 2.20, 0.47, 19, WHITE, True, PP_ALIGN.CENTER, font=FONT_NUM)
    add_text(slide, "頻度はほぼ同等 → 差は“回数”より“余裕”に現れる兆候", 8.40, 5.34, 3.82, 0.55, 12.8, CYAN, True, PP_ALIGN.CENTER)
    add_footer(slide, 14, "Source: 有効離隔43件（拡幅24 / 未拡幅19）｜追越し47件・総動画時間2.48h")

    # 15 — Evidence strength
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=LIGHT, radius=False)
    add_title(slide, "数字は“有望なシグナル”だが、確証にはまだ届かない", "p値だけでなく、効果の方向・区間推定・実件数を並べて判断する", "EVIDENCE STRENGTH")
    evidence_cards = [
        (
            0.62, "組合せで見た方向", f"{superiority*100:.1f}%", "拡幅側の離隔が大きい",
            "拡幅・未拡幅から1件ずつ選ぶ全456組の比較", BLUE,
        ),
        (
            4.48, "中央値差", f"+{widened_median-unwidened_median:.2f} m", "bootstrap 95%区間",
            f"{median_diff_ci[0]:+.2f} ～ {median_diff_ci[1]:+.2f} m（0を含む）", TEAL,
        ),
        (
            8.34, "1.5m未満の実件数", f"{widened_low_n}/{len(widened)}  vs  {unwidened_low_n}/{len(unwidened)}", "拡幅 / 未拡幅",
            f"95%区間 {widened_low_ci[0]*100:.1f}–{widened_low_ci[1]*100:.1f}% / {unwidened_low_ci[0]*100:.1f}–{unwidened_low_ci[1]*100:.1f}%", ORANGE,
        ),
    ]
    for x, label, value, sublabel, note, accent in evidence_cards:
        add_rect(slide, x + 0.05, 1.64, 3.68, 3.18, fill=LINE, radius=True)
        add_rect(slide, x, 1.58, 3.68, 3.18, fill=WHITE, radius=True)
        add_rect(slide, x, 1.58, 3.68, 0.11, fill=accent, radius=False)
        add_text(slide, label, x + 0.24, 1.91, 3.20, 0.34, 13.2, MID, True, PP_ALIGN.CENTER)
        add_text(slide, value, x + 0.18, 2.43, 3.32, 0.72, 27, NAVY, True, PP_ALIGN.CENTER, font=FONT_NUM)
        add_text(slide, sublabel, x + 0.24, 3.17, 3.20, 0.30, 10.5, accent, True, PP_ALIGN.CENTER)
        add_rect(slide, x + 0.23, 3.72, 3.22, 0.67, fill=LIGHT, radius=True)
        add_text(slide, note, x + 0.38, 3.79, 2.92, 0.50, 9.4, MID, True, PP_ALIGN.CENTER)

    add_rect(slide, 0.62, 5.22, 12.08, 1.10, fill=NAVY, radius=True)
    add_text(slide, "現在地", 0.92, 5.43, 0.88, 0.26, 10.5, CYAN, True)
    levels = [
        ("観測", "データ化", GREEN),
        ("シグナル", "今回", BLUE),
        ("再現", "条件調整", MUTED),
        ("因果", "対照設計", MUTED),
    ]
    lx = 1.62
    for i, (name, note, accent) in enumerate(levels):
        add_circle(slide, lx, 5.41, 0.34, accent, str(i + 1), WHITE, 9)
        add_text(slide, name, lx + 0.42, 5.35, 0.62, 0.26, 11.2, WHITE, True)
        add_text(slide, note, lx + 0.42, 5.66, 0.72, 0.20, 8.3, color("B9CADA"))
        if i < len(levels) - 1:
            add_text(slide, "→", lx + 1.12, 5.41, 0.28, 0.30, 13, CYAN if i == 0 else MUTED, True, PP_ALIGN.CENTER)
        lx += 1.48
    add_rect(slide, 7.72, 5.39, 0.02, 0.56, fill=color("355777"), radius=False)
    add_text(slide, "判断：改善候補を絞る根拠には使える。\n効果を断定する根拠にはまだ使わない。", 7.96, 5.30, 4.05, 0.70, 11.2, CYAN, True, PP_ALIGN.CENTER)
    add_footer(slide, 15, "Source: 有効離隔43件｜20,000回bootstrap（seed固定）｜Wilson 95%区間")

    # 16 — Trust and limitations
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=LIGHT, radius=False)
    add_title(slide, "“見える化”と同時に、データの限界も見える化", "運用完了率は高いが、現行DBだけで速度・TTCや因果を結論づけることはできない", "TRUST & LIMITS")
    add_rect(slide, 0.62, 1.56, 4.02, 4.74, fill=WHITE, radius=True)
    add_text(slide, "運用の健全性", 0.95, 1.86, 2.5, 0.32, 15, NAVY, True)
    completion = stats.completed / stats.runs
    add_text(slide, f"{completion*100:.1f}%", 0.95, 2.34, 2.0, 0.65, 30, GREEN, True, font=FONT_NUM)
    add_text(slide, "Run完了", 2.80, 2.55, 0.95, 0.26, 11, MID, True)
    add_rect(slide, 0.95, 3.18, 3.26, 0.22, fill=color("E8EEF4"), radius=True)
    add_rect(slide, 0.95, 3.18, 3.26 * completion, 0.22, fill=GREEN, radius=True)
    add_multiline(slide, [
        f"完了 {stats.completed} / {stats.runs} Run",
        f"処理中 {stats.processing}・エラー {stats.errors}",
        "confidenceは正解率ではない",
    ], 0.95, 3.72, 3.2, 1.28, 12, INK, bullet=True, gap=7)
    add_rect(slide, 0.95, 5.34, 3.22, 0.55, fill=PALE_TEAL, radius=True)
    add_text(slide, "人のレビューで最終確定", 1.12, 5.43, 2.90, 0.30, 12, TEAL, True, PP_ALIGN.CENTER)

    add_rect(slide, 4.92, 1.56, 3.67, 4.74, fill=WHITE, radius=True)
    add_text(slide, "現在の分析制約", 5.25, 1.86, 2.5, 0.32, 15, NAVY, True)
    constraints = [
        ("43件", "有効離隔", ORANGE),
        ("4件", "離隔欠測", RED),
        ("p=.065", "5%水準未達", PURPLE),
        ("1方向", "離隔比較の方向", BLUE),
    ]
    for i, (value, label, accent) in enumerate(constraints):
        col, row = i % 2, i // 2
        xx, yy = 5.24 + col * 1.56, 2.42 + row * 1.28
        add_rect(slide, xx, yy, 1.38, 1.02, fill=LIGHT, radius=True)
        add_text(slide, value, xx + 0.08, yy + 0.15, 1.22, 0.39, 18, accent, True, PP_ALIGN.CENTER, font=FONT_NUM)
        add_text(slide, label, xx + 0.08, yy + 0.59, 1.22, 0.22, 8.8, MID, True, PP_ALIGN.CENTER)
    add_text(slide, "平均差は外れ値の影響を受けるため、中央値と構成比を併記。", 5.24, 5.20, 3.00, 0.64, 10.7, MID, False, PP_ALIGN.CENTER)

    add_rect(slide, 8.87, 1.56, 3.83, 4.74, fill=WHITE, radius=True)
    add_text(slide, "未評価の指標", 9.20, 1.86, 2.5, 0.32, 15, NAVY, True)
    positive_speed = stats.speed_nonnull - stats.speed_zero
    positive_ttc = stats.ttc_nonnull - stats.ttc_zero
    items = [
        ("速度", positive_speed, stats.speed_nonnull, RED),
        ("TTC", positive_ttc, stats.ttc_nonnull, ORANGE),
    ]
    for i, (label, positive, total, accent) in enumerate(items):
        yy = 2.52 + i * 1.20
        add_text(slide, label, 9.20, yy, 0.78, 0.28, 12, NAVY, True)
        add_text(slide, f"正値 {positive:,}", 10.18, yy - 0.02, 1.92, 0.34, 17, accent, True, PP_ALIGN.RIGHT, font=FONT_NUM)
        add_text(slide, f"/ 非NULL {total:,}", 9.20, yy + 0.39, 2.92, 0.24, 9.5, MUTED, False, PP_ALIGN.RIGHT)
        add_rect(slide, 9.20, yy + 0.75, 2.92, 0.12, fill=PALE_RED, radius=False)
    add_rect(slide, 9.20, 5.28, 3.02, 0.64, fill=PALE_RED, radius=True)
    add_text(slide, "現行DBでは実績グラフに使用しない", 9.38, 5.37, 2.66, 0.38, 10.5, RED, True, PP_ALIGN.CENTER)
    add_footer(slide, 16, "Source: 現行DBのNULL/0値監査｜速度482,679件・TTC79,072件はいずれも正値0")

    # 17 — Next steps
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY, radius=False)
    add_title(slide, "次の一手：傾向を、道路改善の判断材料へ", "精度検証と比較設計を整え、地点・時間帯別のリスク可視化へ展開", "NEXT ACTIONS", dark=True)
    actions = [
        ("01", "計測品質を復旧", "速度・TTCパイプラインを検証し、\n0値の原因と再計算手順を確立", RED),
        ("02", "Ground Truthを強化", "47イベントを二重レビュー。\n欠測4件と評価者間一致を管理", ORANGE),
        ("03", "比較条件を揃える", "地点・方向・時間・天候をマッチし、\nサンプル数を拡大", TEAL),
        ("04", "意思決定へ接続", "1.5m未満の構成比などを\n地点別優先順位に変換", BLUE),
    ]
    for i, (num, title, note, accent) in enumerate(actions):
        x = 0.64 + i * 3.08
        add_rect(slide, x, 1.78, 2.76, 3.43, fill=WHITE, radius=True)
        add_circle(slide, x + 0.24, 2.05, 0.58, accent, num, WHITE, 12)
        add_text(slide, title, x + 0.24, 2.84, 2.28, 0.52, 17, NAVY, True, valign=MSO_ANCHOR.TOP)
        add_text(slide, note, x + 0.24, 3.52, 2.28, 1.08, 11.5, MID, False, valign=MSO_ANCHOR.TOP)
        add_rect(slide, x + 0.24, 4.72, 2.28, 0.10, fill=accent, radius=False)
    add_rect(slide, 0.64, 5.67, 12.00, 0.80, fill=BLUE, radius=True)
    add_text(slide, "件数から安全余裕へ —— 映像を、道路改善の定量根拠に。", 0.92, 5.82, 11.44, 0.42, 20, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide, 17, "ADC System｜実データ可視化 可視化拡充版 2026-07-30", dark=True)

    return prs


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stats = load_stats()
    assets = prepare_assets()
    prs = make_presentation(stats, assets)
    prs.save(OUT_PATH)
    print(f"saved: {OUT_PATH}")
    print(f"slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
