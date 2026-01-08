import openpyxl
from openpyxl import load_workbook
import os

# Load the Excel file
excel_path = r'C:\Users\kurok\Downloads\G_ADC\ADC_08\テスト＿卒論.xlsx'

try:
    wb = load_workbook(excel_path)
    print(f"Successfully loaded: {excel_path}")
    print(f"\nAvailable sheets: {wb.sheetnames}")
    
    # Examine each sheet
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n{'='*60}")
        print(f"Sheet: {sheet_name}")
        print(f"{'='*60}")
        print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")
        
        # Display the content
        print("\nContent:")
        for row in ws.iter_rows(min_row=1, max_row=min(20, ws.max_row), values_only=False):
            row_data = []
            for cell in row:
                value = cell.value
                if value is None:
                    row_data.append("[EMPTY]")
                else:
                    row_data.append(str(value))
            print(" | ".join(row_data))
            
except Exception as e:
    print(f"Error: {e}")
    print("\nTrying to install openpyxl...")
    import subprocess
    subprocess.run(["pip", "install", "openpyxl"])
